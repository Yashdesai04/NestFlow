import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def create_nestflow_docx():
    doc = docx.Document()

    # Configure Margins (1 inch all around)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # 1. Body Text Style = 12pt
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)

    # 2. Heading 1 Style = 14pt Bold
    h1 = doc.styles['Heading 1']
    h1.font.name = 'Calibri'
    h1.font.size = Pt(14)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(4)

    # 3. Heading 2 Style = 14pt Bold
    h2 = doc.styles['Heading 2']
    h2.font.name = 'Calibri'
    h2.font.size = Pt(14)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(0xC2, 0x41, 0x0C)
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(3)

    # Cover Page
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(36)
    t1 = title_p.add_run("NESTFLOW: SMART CO-LIVING & PROPERTY MANAGEMENT SYSTEM\n\n")
    t1.font.size = Pt(18)
    t1.font.bold = True
    t1.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s1 = sub_p.add_run("A Project Report\nSubmitted in partial fulfilment for the award of the degree of\n")
    s1.font.size = Pt(13)
    s2 = sub_p.add_run("INTEGRATED MASTER OF SCIENCE - (CA & IT)\n\n")
    s2.font.size = Pt(14)
    s2.font.bold = True

    inst_p = doc.add_paragraph()
    inst_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    inst_p.paragraph_format.space_before = Pt(80)
    i1 = inst_p.add_run("Department of Computer Science\nINDUS UNIVERSITY\nAcademic Year 2025–2026\n")
    i1.font.size = Pt(13)
    i1.font.bold = True
    doc.add_page_break()

    # Certificate
    doc.add_heading("CERTIFICATE", level=1)
    cert_p = doc.add_paragraph()
    cert_p.add_run(
        'This is to certify that the project titled "NESTFLOW: SMART CO-LIVING & PROPERTY MANAGEMENT SYSTEM" '
        'is the bona fide work carried out in partial fulfilment of the requirements for the award of the degree '
        'of Integrated Master of Science (CA & IT) at Department of Computer Science, Indus University '
        'during the academic year 2025–2026.'
    )
    cert_sig = doc.add_paragraph()
    cert_sig.paragraph_format.space_before = Pt(60)
    cert_sig.add_run("Internal Guide Signature: _______________________      Date: ______________\n\n")
    cert_sig.add_run("Head of Dept. Signature:  _______________________      Date: ______________")
    doc.add_page_break()

    # Table of Contents
    doc.add_heading("Table of Contents", level=1)
    toc_data = [
        ("1", "Introduction", "1"),
        ("2", "Organization Profile", "2"),
        ("3", "Existing system", "3"),
        ("4", "Problem areas of existing system", "4"),
        ("5", "Need for the new system", "5"),
        ("6", "Proposed System", "6"),
        ("7", "System Flow Diagram", "7"),
        ("8", "UML Diagrams ( Class, Activity, Sequence, Use case)", "8–13"),
        ("9", "Data Dictionary", "14–19"),
        ("10", "User Interface Screens", "20–25"),
        ("11", "Reports", "26–27"),
        ("12", "Test Plan and Test Cases", "28–30"),
        ("13", "Future Enhancement", "31"),
        ("14", "References", "32")
    ]
    t = doc.add_table(rows=1, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, name in enumerate(["Sr. No.", "Contents", "Page No."]):
        t.rows[0].cells[idx].text = name
        set_cell_background(t.rows[0].cells[idx], "F1F5F9")
        t.rows[0].cells[idx].paragraphs[0].runs[0].font.bold = True
        t.rows[0].cells[idx].paragraphs[0].runs[0].font.size = Pt(12)
    for r in toc_data:
        row = t.add_row().cells
        row[0].text, row[1].text, row[2].text = r
        for c in row:
            if c.paragraphs[0].runs:
                c.paragraphs[0].runs[0].font.size = Pt(12)
    doc.add_page_break()

    # 1. Introduction
    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph(
        "NestFlow is a modern, responsive cloud-native property and co-living management system engineered "
        "to optimize operations across student hostels, paying guest (PG) accommodations, and residential rental complexes. "
        "Built using Next.js 16 (Turbopack), PostgreSQL with Prisma ORM 7, Clerk Identity Authentication, and Razorpay Payments, "
        "NestFlow connects tenants, property managers, and administrators through dedicated, synchronized web portals."
    )
    doc.add_heading("Objectives of the System", level=2)
    doc.add_paragraph("• Provide real-time bed-level occupancy tracking (VACANT, OCCUPIED, MAINTENANCE) to eliminate inventory leakage.")
    doc.add_paragraph("• Automate monthly rent invoicing via scheduled serverless cron jobs running on the 1st of every calendar month.")
    doc.add_paragraph("• Enable secure online rent and utility settlement via Razorpay with cryptographic webhook verification.")
    doc.add_paragraph("• Digitize legal leasing agreements through in-browser digital signatures and downloadable agreement records.")
    doc.add_paragraph("• Minimize mess food waste through a daily resident meal opt-out registration system.")
    doc.add_paragraph("• Strengthen premises security using tokenized QR visitor gate passes with entry and exit timestamp logging.")

    # 2. Organization Profile
    doc.add_heading("2. Organization Profile", level=1)
    doc.add_paragraph("• Project Title: NestFlow Co-Living & Property Management System")
    doc.add_paragraph("• Industry Domain: Property Technology (PropTech) / SaaS Web Application")
    doc.add_paragraph("• Target Deployment: Vercel Edge Serverless Architecture & Cloud PostgreSQL Database")
    
    doc.add_heading("Technology Stack Specifications", level=2)
    tt = doc.add_table(rows=1, cols=3)
    tt.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, name in enumerate(["Category", "Technology / Framework", "System Responsibility"]):
        tt.rows[0].cells[idx].text = name
        set_cell_background(tt.rows[0].cells[idx], "F1F5F9")
        tt.rows[0].cells[idx].paragraphs[0].runs[0].font.bold = True
        tt.rows[0].cells[idx].paragraphs[0].runs[0].font.size = Pt(12)

    tech_records = [
        ("Frontend", "Next.js 16.3 (Turbopack), React 19", "Server Components & Client State Hydration"),
        ("Styling", "Tailwind CSS v4, Lucide React", "Responsive Utility-First Design System"),
        ("Backend", "Next.js Server Actions & Route Handlers", "Type-safe Mutations, APIs, and Webhooks"),
        ("Database", "PostgreSQL, Prisma ORM 7", "Relational Storage with @prisma/adapter-pg Driver"),
        ("Authentication", "Clerk Authentication SDK", "JWT Session Management & Webhook Sync"),
        ("Payment Gateway", "Razorpay Standard Checkout & APIs", "Orders API & Cryptographic HMAC Verification"),
        ("Scheduler", "Vercel Cron Engine (vercel.json)", "1st-of-month Automated Invoicing Routine"),
        ("Development Tools", "VS Code, Git/GitHub, Draw.io", "IDE, Version Control, & UML Modeling")
    ]
    for r in tech_records:
        row = tt.add_row().cells
        row[0].text, row[1].text, row[2].text = r
        for c in row:
            c.paragraphs[0].runs[0].font.size = Pt(12)
    doc.add_page_break()

    # 3. Existing system
    doc.add_heading("3. Existing system", level=1)
    doc.add_paragraph(
        "In traditional PG and rental housing environments, property managers rely heavily on physical paper ledgers, "
        "informal messaging groups (such as WhatsApp), and manual payment reconciliations. Key characteristics include:"
    )
    doc.add_paragraph("1. Physical Notebooks & Registers: Room allotments, tenant contact details, and advance deposits are logged by hand.")
    doc.add_paragraph("2. Offline Payment Reconciliation: Tenants pay via direct UPI transfers or cash, sending screenshot receipts that require manual ledger checking.")
    doc.add_paragraph("3. Paper Gate Entry Books: Security staff record visitor details in physical logbooks, often resulting in illegible handwriting and unverified phone numbers.")
    doc.add_paragraph("4. Uncoordinated Kitchen Preparation: Mess kitchens cook fixed quantities of meals daily without knowing how many residents are away, leading to severe food wastage.")

    # 4. Problem areas of existing system
    doc.add_heading("4. Problem areas of existing system", level=1)
    doc.add_paragraph("• Revenue Leakage: Uncollected utility bills and delayed monthly rent collections due to lack of automated billing.")
    doc.add_paragraph("• Double Bookings & Overcrowding: Property owners lack real-time visibility into bed occupancy across shared rooms.")
    doc.add_paragraph("• Untracked Maintenance Requests: Verbal maintenance complaints are frequently forgotten or delayed without accountability.")
    doc.add_paragraph("• Substantial Kitchen Food Waste: Fixed daily cooking without resident meal attendance planning inflates mess operational costs.")
    doc.add_paragraph("• Security Vulnerabilities: Absence of digital visitor identity verification poses safety concerns for residents.")

    # 5. Need for the new system
    doc.add_heading("5. Need for the new system", level=1)
    doc.add_paragraph(
        "NestFlow addresses the vulnerabilities of manual operations through a centralized, automated web platform that provides:"
    )
    doc.add_paragraph("• Single-Pane Management: Dedicated, role-based dashboards for Tenants, Landlords, and Super Administrators.")
    doc.add_paragraph("• Real-Time Bed Availability: Instant bed inventory updates prevent double allocation and ensure unbooked beds are listed immediately.")
    doc.add_paragraph("• Automated Recurring Financials: Automated cron invoicing on the 1st of every month coupled with direct online settlement via Razorpay.")
    doc.add_paragraph("• Digital Accountability: In-browser legal lease execution, verifiable QR visitor gate passes, and structured maintenance tracking.")
    doc.add_page_break()

    # 6. Proposed System
    doc.add_heading("6. Proposed System", level=1)
    doc.add_paragraph("The proposed platform provides specialized modules tailored to each operational stakeholder:")
    doc.add_paragraph(
        "1. Tenant Living Portal: Responsive touch Rotary Hub dial navigation, lease review and digital signing, "
        "online rent invoice settlement via Razorpay modal, meal opt-outs, and visitor gate pass issuance."
    )
    doc.add_paragraph(
        "2. Landlord Operations Desk: Room and bed configuration, visual real-time bed occupancy grid, rental application processing, "
        "mess menu scheduling, maintenance vendor assignments, and atomic resident offboarding."
    )
    doc.add_paragraph(
        "3. Admin Governance Module: Multi-listing moderation queue, tenant and landlord KYC verification desk, "
        "system analytics dashboard, user role elevation (RBAC), and security audit monitoring."
    )

    # 7. System Flow Diagram
    doc.add_heading("7. System Flow Diagram", level=1)
    doc.add_paragraph("The system architecture follows a decoupled serverless execution flow:")
    doc.add_paragraph("1. Authentication: Resident signs in through Clerk Auth -> Profile synced to PostgreSQL via /api/webhooks/clerk.")
    doc.add_paragraph("2. Discovery & Application: Tenant selects bed in property -> Submits digital rental application.")
    doc.add_paragraph("3. Bed Allocation: Landlord accepts application -> Assigns specific bed space (status shifts to OCCUPIED).")
    doc.add_paragraph("4. Legal Signing: Tenant views lease in portal -> Signs on HTML5 canvas pad -> Lease status marked ACTIVE.")
    doc.add_paragraph("5. Automated Invoicing: Cron executes on 1st of month -> Generates pending rent invoices for all active signed leases.")
    doc.add_paragraph("6. Settlement: Tenant opens Razorpay checkout modal -> Settles bill -> Signature checked -> Payment marked PAID.")
    doc.add_paragraph("7. Offboarding: Landlord clicks Vacate Resident -> Database transaction terminates agreement and reverts bed to VACANT.")
    doc.add_page_break()

    # 8. UML Diagrams
    doc.add_heading("8. UML Diagrams ( Class, Activity, Sequence, Use case)", level=1)
    
    doc.add_heading("8.1 Use Case Diagram", level=2)
    doc.add_paragraph("• Super Admin: Log in, Inspect KYC queue, Approve/Reject property listings, Manage users, Inspect logs, Log out.")
    doc.add_paragraph("• Tenant: Register/Login, Browse listings, Submit application, Sign lease, Pay rent, Issue gate pass, Opt out of meals, Raise maintenance ticket, Log out.")
    doc.add_paragraph("• Landlord: Manage properties, Configure beds, Review applications, Assign beds, Set mess menus, Assign staff, Offboard resident.")

    doc.add_heading("8.2 Activity Diagram", level=2)
    doc.add_paragraph(
        "• KYC Moderation: Admin Login -> Open KYC Queue -> Inspect Document -> Set APPROVED / REJECTED -> Save DB -> Refresh Queue.\n"
        "• Online Checkout: Tenant Login -> Invoices -> Pay Rent -> Create Razorpay Order -> Launch SDK Modal -> Authorize -> Validate Signature -> Mark PAID in DB."
    )

    doc.add_heading("8.3 Sequence Diagram", level=2)
    doc.add_paragraph(
        "• Rent Settlement: Tenant -> Client -> API (/api/payments/create-order) -> Razorpay API (Create Order) -> Client (Open Modal) -> Tenant (Authorize) -> Razorpay (Signature) -> Server Action (verifyPaymentAction) -> DB (Set status=PAID)."
    )

    doc.add_heading("8.4 Class Diagram", level=2)
    doc.add_paragraph("• User: id, clerkId, email, fullName, phone, role, kycStatus, isVerified")
    doc.add_paragraph("• Property: id, landlordId, title, propertyType, leasingType, priceMonthly, isVerified")
    doc.add_paragraph("• Room: id, propertyId, roomNumber, roomType")
    doc.add_paragraph("• Bed: id, roomId, bedNumber, priceMonthly, status (VACANT, OCCUPIED, MAINTENANCE)")
    doc.add_paragraph("• Lease: id, tenantId, propertyId, roomId, bedId, startDate, endDate, rentAmount, isSigned, status")
    doc.add_paragraph("• Payment: id, leaseId, tenantId, rentAmount, utilityAmount, totalAmount, status (PENDING, PAID)")
    doc.add_paragraph("• GatePass: id, propertyId, tenantId, visitorName, passCode, entryTime, exitTime, status")
    doc.add_paragraph("• MaintenanceTicket: id, propertyId, tenantId, title, priority, status, assignedStaff")
    doc.add_page_break()

    # 9. Data Dictionary
    doc.add_heading("9. Data Dictionary", level=1)
    dd_table = doc.add_table(rows=1, cols=4)
    dd_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, name in enumerate(["Field Name", "Data Type", "Constraints", "Description"]):
        dd_table.rows[0].cells[idx].text = name
        set_cell_background(dd_table.rows[0].cells[idx], "F1F5F9")
        dd_table.rows[0].cells[idx].paragraphs[0].runs[0].font.bold = True
        dd_table.rows[0].cells[idx].paragraphs[0].runs[0].font.size = Pt(12)

    dd_data = [
        ("User.id", "VARCHAR(30)", "PRIMARY KEY (CUID)", "Unique internal user identifier"),
        ("User.clerkId", "VARCHAR(255)", "UNIQUE, NOT NULL", "Clerk Authentication Identity"),
        ("User.role", "ENUM", "DEFAULT 'TENANT'", "TENANT, LANDLORD, ADMIN"),
        ("User.kycStatus", "ENUM", "DEFAULT 'UNSUBMITTED'", "UNSUBMITTED, PENDING, APPROVED, REJECTED"),
        ("Property.id", "VARCHAR(30)", "PRIMARY KEY (CUID)", "Unique property identifier"),
        ("Property.landlordId", "VARCHAR(30)", "FOREIGN KEY -> User(id)", "Associated landlord owner ID"),
        ("Property.leasingType", "ENUM", "DEFAULT 'BED_LEVEL'", "BED_LEVEL or FULL_UNIT"),
        ("Bed.status", "ENUM", "DEFAULT 'VACANT'", "VACANT, OCCUPIED, MAINTENANCE"),
        ("Lease.rentAmount", "FLOAT", "NOT NULL", "Monthly agreed rent amount (INR)"),
        ("Lease.isSigned", "BOOLEAN", "DEFAULT false", "Digital signature executed flag"),
        ("Payment.totalAmount", "FLOAT", "NOT NULL", "Sum of base rent + split utilities"),
        ("Payment.status", "ENUM", "DEFAULT 'PENDING'", "PENDING, PAID, OVERDUE, FAILED"),
        ("GatePass.passCode", "VARCHAR(50)", "NOT NULL", "4-digit security entry token (VP-XXXX)"),
        ("GatePass.status", "ENUM", "DEFAULT 'APPROVED'", "APPROVED, CHECKED_IN, EXPIRED")
    ]
    for r in dd_data:
        row = dd_table.add_row().cells
        row[0].text, row[1].text, row[2].text, row[3].text = r
        for c in row:
            c.paragraphs[0].runs[0].font.size = Pt(12)
    doc.add_page_break()

    # 10. User Interface Screens
    doc.add_heading("10. User Interface Screens", level=1)
    doc.add_paragraph("• UI-01: Public Property Catalog (/properties): Search and filter by gender policy, room type, and price with photo galleries.")
    doc.add_paragraph("• UI-02: Tenant Rotary Command Hub (/dashboard/tenant): Circular touch-enabled dial menu providing one-click resident navigation.")
    doc.add_paragraph("• UI-03: Digital Lease E-Sign Desk (/dashboard/tenant/lease): In-browser agreement viewer featuring legal terms and signature pad.")
    doc.add_paragraph("• UI-04: Tenant Payments & Razorpay Modal (/dashboard/tenant/payments): Invoice ledger displaying due amounts with direct checkout.")
    doc.add_paragraph("• UI-05: Landlord Visual Bed Matrix (/dashboard/landlord/bed-grid): Room-by-room grid with live color-coded bed occupancy tags.")
    doc.add_paragraph("• UI-06: Visitor Gate Pass Desk (/dashboard/tenant/gatepass): Two-column layout for guest token generation and security check-ins.")
    doc.add_paragraph("• UI-07: Landlord Resident Directory (/dashboard/landlord/tenants): Occupancy roster with one-click Vacate & Free Bed modal.")
    doc.add_paragraph("• UI-08: Admin Moderation Center (/dashboard/admin): Executive queues for inspecting KYC proof and approving property listings.")

    # 11. Reports
    doc.add_heading("11. Reports", level=1)
    doc.add_paragraph("• Monthly Rent Collection & Arrears Ledger: Summarizes total billed revenue, collected payments, and overdue accounts per property.")
    doc.add_paragraph("• Bed Occupancy & Inventory Yield Report: Tracks real-time occupancy rates, vacant bed durations, and turnover statistics.")
    doc.add_paragraph("• Visitor Security Audit Report: Chronological audit of issued visitor gate passes, check-in timestamps, and departure records.")
    doc.add_paragraph("• Maintenance Resolution & Cost Report: Aggregates ticket turnaround times, assigned vendors, and cumulative repair expenditure.")
    doc.add_paragraph("• Mess Kitchen Attendance Projection: Daily forecast calculating required meal counts based on resident opt-out submissions.")
    doc.add_page_break()

    # 12. Test Plan and Test Cases
    doc.add_heading("12. Test Plan and Test Cases", level=1)
    test_table = doc.add_table(rows=1, cols=4)
    test_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, name in enumerate(["Test ID", "Scenario", "Expected Result", "Status"]):
        test_table.rows[0].cells[idx].text = name
        set_cell_background(test_table.rows[0].cells[idx], "F1F5F9")
        test_table.rows[0].cells[idx].paragraphs[0].runs[0].font.bold = True
        test_table.rows[0].cells[idx].paragraphs[0].runs[0].font.size = Pt(12)

    tests = [
        ("TC01", "Unauthenticated route access", "Redirect to /login", "Pass"),
        ("TC02", "Clerk identity webhook sync", "User created in PostgreSQL", "Pass"),
        ("TC03", "Double-assigning occupied bed", "Throws Occupied validation error", "Pass"),
        ("TC04", "Razorpay order creation", "Returns valid order ID and amount", "Pass"),
        ("TC05", "HMAC signature verification", "Payment record marked PAID", "Pass"),
        ("TC06", "Signature tampering protection", "Server rejects payment; remains PENDING", "Pass"),
        ("TC07", "Digital lease canvas submission", "Lease status becomes ACTIVE", "Pass"),
        ("TC08", "Monthly cron authentication", "Unauthorized call returns 401", "Pass"),
        ("TC09", "Invoicing cron idempotence", "Skips existing bills; creates 0 duplicates", "Pass"),
        ("TC10", "Visitor gate pass token generation", "Unique VP-XXXX token generated", "Pass"),
        ("TC11", "Daily meal opt-out skipping", "Decrements kitchen diner count", "Pass"),
        ("TC12", "Resident offboarding & vacate", "Lease TERMINATED; bed becomes VACANT", "Pass")
    ]
    for r in tests:
        row = test_table.add_row().cells
        row[0].text, row[1].text, row[2].text, row[3].text = r
        for c in row:
            c.paragraphs[0].runs[0].font.size = Pt(12)
    doc.add_page_break()

    # 13. Future Enhancement
    doc.add_heading("13. Future Enhancement", level=1)
    doc.add_paragraph("1. Mobile Application (React Native / Flutter): Native mobile push alerts, biometric login, and offline QR visitor passes.")
    doc.add_paragraph("2. IoT Smart Lock Keyless Entry: Automatic time-limited keypad PINs synchronized with approved gate passes and active leases.")
    doc.add_paragraph("3. Automated Utility Sub-Metering: Direct integration with digital room electric meters to split monthly kilowatt-hour consumption.")
    doc.add_paragraph("4. AI Mess Food Forecasting: Machine learning models predicting meal attendance to minimize kitchen food waste.")
    doc.add_paragraph("5. WhatsApp Business Notification Bot: Automated messaging for rent reminders, receipts, and visitor security codes.")

    # 14. References
    doc.add_heading("14. References", level=1)
    doc.add_paragraph("• Next.js 16 Documentation - https://nextjs.org/docs")
    doc.add_paragraph("• Prisma ORM 7 Guide & Driver Adapters - https://www.prisma.io/docs")
    doc.add_paragraph("• Clerk Authentication SDK - https://clerk.com/docs")
    doc.add_paragraph("• Razorpay Developers Guide & Webhooks - https://razorpay.com/docs")
    doc.add_paragraph("• PostgreSQL Documentation - https://www.postgresql.org/docs")
    doc.add_paragraph("• Tailwind CSS v4 Design Reference - https://tailwindcss.com/docs")

    doc.save("NestFlow_Project_Report.docx")
    print("SUCCESS: 'NestFlow_Project_Report.docx' has been successfully created!")

if __name__ == "__main__":
    create_nestflow_docx()